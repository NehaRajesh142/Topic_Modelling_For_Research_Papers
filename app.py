
from flask import Flask, render_template, request, jsonify, send_file
import csv, os, re, math, random, io, urllib.parse, urllib.request, json
from collections import Counter

app=Flask(__name__)
DATA_FILE="research_papers.csv"
EXCEL_FILE="research_papers.xlsx"

STOPWORDS=set("""a an the and or but if then than so because as at by for from in into on of to with
without is are was were be been being this that these those it its they them their we our you your
i he she his her has have had do does did can could will would should may might must about above
after again against all am any around away before below between both during each few further here
how more most other over own same some such through under until very what when where which while
who whom why not no nor only just also paper papers study studies research using use used based
approach method methods results show shows shown new proposed provide provides""".split())

def read_rows():
    if not os.path.exists(DATA_FILE): return []
    with open(DATA_FILE,encoding="utf-8-sig",newline="") as f: return list(csv.DictReader(f))

def save_rows(rows):
    with open(DATA_FILE,"w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=["title","abstract","year","author","source","doi"])
        w.writeheader()
        for r in rows:
            w.writerow({k:r.get(k,"") for k in w.fieldnames})

def tokens(text):
    ws=re.findall(r"[a-zA-Z]{3,}",(text or "").lower())
    return [w for w in ws if w not in STOPWORDS]

def lda(docs,k=5,iters=100):
    vocab=sorted(set(w for d in docs for w in d))
    if len(vocab)<k*3: vocab=sorted(Counter(w for d in docs for w in d),key=lambda x:Counter(w for d in docs for w in d)[x],reverse=True)
    vi={w:i for i,w in enumerate(vocab)}
    nw=[[0]*len(vocab) for _ in range(k)]; nd=[[0]*k for _ in docs]; nk=[0]*k; z=[]
    random.seed(42); alpha=.5; beta=.1
    for di,d in enumerate(docs):
        zz=[]
        for w in d:
            if w not in vi: continue
            t=random.randrange(k); zz.append(t); nw[t][vi[w]]+=1; nd[di][t]+=1; nk[t]+=1
        z.append(zz)
    for _ in range(iters):
        for di,d in enumerate(docs):
            pos=0
            for w in d:
                if w not in vi: continue
                old=z[di][pos]; idx=vi[w]
                nw[old][idx]-=1; nd[di][old]-=1; nk[old]-=1
                probs=[(nd[di][t]+alpha)*(nw[t][idx]+beta)/(nk[t]+len(vocab)*beta) for t in range(k)]
                r=random.random()*sum(probs); acc=0; new=0
                for t,p in enumerate(probs):
                    acc+=p
                    if r<=acc: new=t; break
                z[di][pos]=new; nw[new][idx]+=1; nd[di][new]+=1; nk[new]+=1; pos+=1
    topic_words={}
    for t in range(k):
        inds=sorted(range(len(vocab)),key=lambda i:nw[t][i],reverse=True)
        topic_words[t+1]=[vocab[i] for i in inds[:10] if nw[t][i]>0]
    probs=[]
    for di in range(len(docs)):
        total=sum(nd[di]); probs.append([nd[di][t]/total if total else 0 for t in range(k)])
    return topic_words,probs

def quality(words,docs):
    if len(words)<2:return 0
    sets=[set(d) for d in docs]; vals=[]
    for i in range(len(words)):
        for j in range(i+1,len(words)):
            a,b=words[i],words[j]; da=sum(a in s for s in sets); dab=sum(a in s and b in s for s in sets)
            vals.append(math.log((dab+1)/(da+1)))
    raw=sum(vals)/len(vals)
    return round(max(0,min(100,50+raw*35)),1)


def topic_name(words):
    w=set(x.lower() for x in words)
    rules=[
        ({"image","classification","vision","detection","segmentation","cnn","convolutional","object"},"Computer Vision & Image Processing"),
        ({"language","text","nlp","transformer","sentiment","semantic","linguistic"},"Natural Language Processing"),
        ({"network","intrusion","cybersecurity","security","threat","encryption","attack"},"Cybersecurity & Network Security"),
        ({"medical","disease","healthcare","diagnosis","heart","clinical","patient"},"Healthcare & Medical AI"),
        ({"solar","renewable","energy","grid","forecasting","power","electricity"},"Renewable Energy & Smart Grids"),
        ({"autonomous","vehicle","road","driving","traffic","object"},"Autonomous Vehicles"),
        ({"cloud","computing","resource","allocation","server","iot","internet","things"},"Cloud Computing & IoT"),
        ({"big","data","analytics","smart","cities","urban"},"Big Data & Smart Cities"),
        ({"blockchain","secure","sharing","ledger","cryptocurrency"},"Blockchain & Secure Data Sharing"),
        ({"machine","learning","algorithm","prediction","model","classification"},"Machine Learning & Prediction"),
    ]
    best_name="Research & Machine Learning"
    best_score=0
    for keys,name in rules:
        score=len(w & keys)
        if score>best_score:
            best_score=score; best_name=name
    return best_name

def model(k):
    rows=read_rows(); docs=[tokens((r.get("title","")+" "+r.get("abstract",""))) for r in rows]
    tw,pr=lda(docs,k)
    out=[]
    for i,r in enumerate(rows):
        t=max(range(k),key=lambda x:pr[i][x])+1
        out.append({**r,"topic":t,"confidence":round(pr[i][t-1]*100,1)})
    years=sorted({int(r["year"]) for r in rows if str(r.get("year","")).isdigit()})
    heat={str(t):[sum(1 for r in out if r["topic"]==t and int(r["year"])==y) for y in years] for t in range(1,k+1)}
    qs={str(t):quality(tw[t],docs) for t in range(1,k+1)}
    names={str(t):topic_name(tw[t]) for t in range(1,k+1)}
    return {"papers":out,"topics":tw,"topic_names":names,"quality":qs,"overall":round(sum(qs.values())/len(qs),1),
            "years":years,"heat":heat,"counts":{str(t):sum(1 for r in out if r["topic"]==t) for t in range(1,k+1)}}

@app.route("/")
def home(): return render_template("index.html",count=len(read_rows()))

@app.get("/api/model")
def api_model():
    return jsonify(model(max(2,min(8,int(request.args.get("topics",5))))))

@app.post("/api/add")
def add():
    d=request.get_json() or {}
    if not d.get("title") or not d.get("abstract") or not d.get("year"): return jsonify(error="Title, abstract and year are required."),400
    rows=read_rows()
    rows.append({"title":d["title"].strip(),"abstract":d["abstract"].strip(),"year":str(d["year"]).strip(),
                 "author":d.get("author","").strip(),"source":d.get("source","Manual entry").strip(),"doi":d.get("doi","").strip()})
    save_rows(rows); return jsonify(success=True,count=len(rows))

@app.get("/api/search-web")
def search_web():
    q=request.args.get("q","").strip()
    if not q:return jsonify(error="Enter a research topic or paper title."),400
    url="https://api.crossref.org/works?query.bibliographic="+urllib.parse.quote(q)+"&rows=8&select=DOI,title,author,published,abstract,container-title"
    try:
        req=urllib.request.Request(url,headers={"User-Agent":"ResearchScope/1.0 (educational project)"})
        data=json.loads(urllib.request.urlopen(req,timeout=12).read().decode("utf-8"))
        results=[]
        for item in data["message"]["items"]:
            title=(item.get("title") or [""])[0]
            abstract=re.sub("<[^>]+>"," ",item.get("abstract","") or "").strip()
            if not abstract: abstract="Abstract not supplied by the metadata provider."
            year=(item.get("published",{}).get("date-parts") or [[None]])[0][0]
            authors=", ".join(((a.get("given","")+" "+a.get("family","")).strip()) for a in item.get("author",[])[:4])
            results.append({"title":title,"abstract":abstract,"year":year or "","author":authors,
                            "doi":item.get("DOI",""),"source":"Crossref"})
        return jsonify(results=results)
    except Exception as e:
        return jsonify(error="Web search failed. Check your internet connection. Details: "+str(e)),502

@app.post("/api/add-web")
def add_web():
    d=request.get_json() or {}; rows=read_rows()
    existing={r.get("doi","").lower() for r in rows if r.get("doi")}
    if d.get("doi") and d["doi"].lower() in existing:return jsonify(error="This DOI is already in your dataset."),400
    rows.append({"title":d.get("title",""),"abstract":d.get("abstract",""),"year":str(d.get("year","")),
                 "author":d.get("author",""),"source":"Crossref","doi":d.get("doi","")})
    save_rows(rows); return jsonify(success=True,count=len(rows))

@app.get("/download/excel")
def excel():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except Exception:
        return jsonify(error="Install openpyxl with: python -m pip install openpyxl"),500
    rows=read_rows()
    wb=Workbook(); ws=wb.active; ws.title="Research Papers"
    headers=["title","abstract","year","author","source","doi"]
    for c,h in enumerate(headers,1):
        cell=ws.cell(1,c,h); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="245E55")
    for r in rows: ws.append([r.get(h,"") for h in headers])
    for row in ws.iter_rows():
        for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
    widths=[42,80,12,28,16,34]
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    wb.save(EXCEL_FILE)
    return send_file(EXCEL_FILE,as_attachment=True,download_name="research_papers.xlsx")

if __name__=="__main__": app.run(debug=True)
